import os
import tempfile
from typing import Any
from module_1_document_processing.composio_connector.events.canonical_event import CanonicalEvent
from module_1_document_processing.parsing.parsed_document import ParsedDocument

try:
    from llama_parse import LlamaParse
except ImportError:
    LlamaParse = None

class LlamaParserService:
    """Cloud OCR & Multimodal Parser supporting PDF, Word, Excel, PowerPoint via LlamaParse with local fallback."""

    def __init__(self) -> None:
        self.api_key = os.environ.get("LLAMAPARSE_API_KEY", "") or os.environ.get("LLAMA_CLOUD_API_KEY", "")
        self.parser = None
        if LlamaParse is not None and self.api_key:
            try:
                self.parser = LlamaParse(api_key=self.api_key, result_type="markdown")
                print("[LlamaParserService] Initialized LlamaParse client with API key.")
            except Exception as err:
                print(f"[LlamaParserService] Error initializing LlamaParse client: {err}")

    def parse_attachment_bytes(self, filename: str, mime_type: str, raw_bytes: bytes | None) -> tuple[str, str, str]:
        """
        Parses attachment bytes using LlamaParse or local fallback.
        Returns (parsed_markdown, parser_used, parse_status).
        """
        # Never return invented text: some callers ignore the status and use the
        # string directly, so a placeholder would be chunked and indexed as if it
        # were real document content.
        if raw_bytes is None or len(raw_bytes) == 0:
            return "", "none", "SKIPPED"

        # Attachments are untrusted input too (arch.md: "Composio output =
        # untrusted input"), so they run the same size and malware checks as an
        # uploaded file rather than being parsed straight away.
        from module_1_document_processing.pipeline import ingestion_guards as _guards

        scan = _guards.security_scanner.scan_raw_bytes(raw_bytes)
        if not scan.is_safe:
            print(f"[LlamaParserService] Rejected attachment '{filename}': {scan.reason}")
            return "", "rejected", "FAILED"

        # Audio/video must be parked, not decoded into noise. These handlers are
        # reached before the MIME router for Gmail/Slack, so route here too.
        from module_1_document_processing.parsing.media_queue import MEDIA_PENDING
        from module_1_document_processing.parsing.mime_router import MIMERouter, ParserCategory

        if MIMERouter().route(mime_type=mime_type, file_path=filename) in (
            ParserCategory.AUDIO,
            ParserCategory.VIDEO,
        ):
            print(f"[LlamaParserService] Attachment '{filename}' is media; parked pending transcription.")
            return "", "media_pending", MEDIA_PENDING

        if self.parser is not None:
            try:
                suffix = f"_{filename}" if filename else ".pdf"
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                    tmp.write(raw_bytes)
                    tmp_path = tmp.name

                documents = self.parser.load_data(tmp_path)
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass

                parsed_text = "\n\n".join([doc.text for doc in documents if hasattr(doc, "text") and doc.text])
                if parsed_text.strip():
                    return parsed_text, "llama_parse", "SUCCESS"
            except Exception as err:
                print(f"[LlamaParserService] LlamaParse cloud parsing error for {filename}: {err}")

        # Fallback for plain text / basic formats
        if mime_type.startswith("text/") or filename.endswith((".txt", ".csv", ".json", ".md")):
            try:
                return raw_bytes.decode("utf-8", errors="replace"), "local_text", "SUCCESS"
            except Exception:
                pass

        # Fast local PDF extraction via pypdf
        if mime_type == "application/pdf" or filename.lower().endswith(".pdf"):
            try:
                import io
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
                pages_text = [page.extract_text() or "" for page in reader.pages]
                pdf_text = "\n\n".join([p.strip() for p in pages_text if p.strip()])
                if pdf_text.strip():
                    return pdf_text, "pypdf_local", "SUCCESS"
            except Exception as pdf_err:
                print(f"[LlamaParserService] Local pypdf error for {filename}: {pdf_err}")

        # Fast local DOCX extraction via python-docx
        if "wordprocessingml" in mime_type or filename.lower().endswith(".docx"):
            try:
                import io
                import docx
                doc = docx.Document(io.BytesIO(raw_bytes))
                docx_text = "\n\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
                if docx_text.strip():
                    return docx_text, "docx_local", "SUCCESS"
            except Exception as docx_err:
                print(f"[LlamaParserService] Local docx error for {filename}: {docx_err}")

        # Nothing could be extracted. Report the failure instead of fabricating
        # content that would be embedded and returned as a search result.
        print(f"[LlamaParserService] No text could be extracted from attachment '{filename}'.")
        return "", "unsupported", "FAILED"

    def parse(self, event: CanonicalEvent, raw_bytes: bytes | None = None) -> ParsedDocument:
        doc_id = f"{event.tenant_id}:{event.source}:{event.external_id}"
        filename = event.metadata.get("name") or event.metadata.get("title") or f"{event.external_id}.pdf"
        mime_type = event.metadata.get("mime_type", "application/pdf")

        # Use LlamaParse if key & raw bytes are available
        if self.parser is not None and raw_bytes is not None:
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=f"_{filename}") as tmp:
                    tmp.write(raw_bytes)
                    tmp_path = tmp.name

                documents = self.parser.load_data(tmp_path)
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass

                parsed_text = "\n\n".join([doc.text for doc in documents if hasattr(doc, "text")])
                return ParsedDocument(
                    doc_id=doc_id,
                    tenant_id=event.tenant_id,
                    user_id=event.user_id,
                    source=event.source,
                    external_id=event.external_id,
                    acl=list(event.acl),
                    mime_type=mime_type,
                    text_content=parsed_text,
                    parse_status="SUCCESS",
                    parser_used="llama_parse",
                    metadata=event.metadata,
                )
            except Exception as err:
                print(f"[LlamaParserService] LlamaParse cloud parsing error: {err}. Falling back to local parser.")

        # Local fallback parser. Content the connector already supplied (Notion
        # blocks, Slack text) is legitimate; inventing filler text is not, so a
        # document with nothing extractable is reported as a parse failure rather
        # than indexed as though it had content.
        fallback_text = (
            event.metadata.get("text_content")
            or event.metadata.get("body")
            or event.metadata.get("text")
        )

        if not fallback_text or not str(fallback_text).strip():
            print(f"[LlamaParserService] No extractable content for {filename}; reporting parse failure.")
            return ParsedDocument(
                doc_id=doc_id,
                tenant_id=event.tenant_id,
                user_id=event.user_id,
                source=event.source,
                external_id=event.external_id,
                acl=list(event.acl),
                mime_type=mime_type,
                text_content="",
                parse_status="FAILED",
                parser_used="fallback_local",
                metadata={**(event.metadata or {}), "error": "No document content could be extracted"},
            )

        return ParsedDocument(
            doc_id=doc_id,
            tenant_id=event.tenant_id,
            user_id=event.user_id,
            source=event.source,
            external_id=event.external_id,
            acl=list(event.acl),
            mime_type=mime_type,
            text_content=f"# {filename}\n\n{fallback_text}",
            parse_status="SUCCESS",
            parser_used="fallback_local",
            metadata=event.metadata,
        )
