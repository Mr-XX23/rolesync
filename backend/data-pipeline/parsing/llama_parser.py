import os
import tempfile
from typing import Any
from connectors.events.canonical_event import CanonicalEvent
from parsing.parsed_document import ParsedDocument

try:
    from llama_parse import LlamaParse
except ImportError:
    LlamaParse = None

class LlamaParserService:
    """Wrapper around LlamaParse SDK with local fallback readers."""

    def __init__(self) -> None:
        self.api_key = os.environ.get("LLAMAPARSE_API_KEY", "")
        self.parser = None
        if LlamaParse is not None and self.api_key:
            try:
                self.parser = LlamaParse(
                    api_key=self.api_key,
                    result_type="markdown",  # Output format as Markdown
                    verbose=False,
                )
            except Exception as e:
                print(f"[LlamaParserService] Error initializing LlamaParse: {e}")

    def parse(self, event: CanonicalEvent, raw_bytes: bytes | None = None) -> ParsedDocument:
        doc_id = f"{event.tenant_id}:{event.source}:{event.external_id}"
        filename = event.metadata.get("name") or f"doc_{event.external_id}"
        ext = event.metadata.get("ext") or ".pdf"

        # If raw_bytes provided and LlamaParse initialized, use LlamaParse
        if raw_bytes and self.parser is not None:
            try:
                with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                    tmp.write(raw_bytes)
                    tmp_path = tmp.name

                documents = self.parser.load_data(tmp_path)
                os.unlink(tmp_path)

                if documents:
                    full_md = "\n\n--- Page Break ---\n\n".join([doc.text for doc in documents])
                    return ParsedDocument(
                        doc_id=doc_id,
                        tenant_id=event.tenant_id,
                        user_id=event.user_id,
                        source=event.source,
                        external_id=event.external_id,
                        acl=list(event.acl),
                        mime_type=event.metadata.get("mime_type", "application/pdf"),
                        text_content=full_md,
                        metadata=event.metadata,
                        page_count=len(documents),
                        parse_status="SUCCESS",
                        parser_used="llama_parse",
                    )
            except Exception as err:
                print(f"[LlamaParserService] LlamaParse call failed: {err}. Using local fallback parser.")

        # Fallback to local parsing logic if LlamaParse is unavailable or raw_bytes is empty
        return self._local_fallback_parse(event, raw_bytes)

    def _local_fallback_parse(self, event: CanonicalEvent, raw_bytes: bytes | None = None) -> ParsedDocument:
        doc_id = f"{event.tenant_id}:{event.source}:{event.external_id}"
        text_content = ""
        page_count = 1

        if raw_bytes:
            # 1. Try PyPDF fallback for PDFs
            try:
                import io
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
                page_count = len(reader.pages)
                pages_text = [page.extract_text() or "" for page in reader.pages]
                text_content = "\n\n--- Page Break ---\n\n".join(pages_text)
            except Exception:
                pass

            # 2. Try docx fallback if text_content is still empty
            if not text_content:
                try:
                    import io
                    import docx
                    doc = docx.Document(io.BytesIO(raw_bytes))
                    text_content = "\n".join([p.text for p in doc.paragraphs if p.text])
                except Exception:
                    pass

        # 3. Final fallback to metadata text
        if not text_content:
            text_content = (
                event.metadata.get("text")
                or event.metadata.get("subject")
                or event.metadata.get("title")
                or f"# {event.metadata.get('name', 'Document')}\n\nDocument content from {event.source}."
            )

        return ParsedDocument(
            doc_id=doc_id,
            tenant_id=event.tenant_id,
            user_id=event.user_id,
            source=event.source,
            external_id=event.external_id,
            acl=list(event.acl),
            mime_type=event.metadata.get("mime_type", "application/pdf"),
            text_content=text_content,
            metadata=event.metadata,
            page_count=page_count,
            parse_status="SUCCESS",
            parser_used="fallback_local",
        )
